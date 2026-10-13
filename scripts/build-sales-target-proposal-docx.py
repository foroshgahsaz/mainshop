#!/usr/bin/env python3
"""Generate management proposal Word doc (Persian, RTL)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def set_paragraph_rtl(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    p_pr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    set_paragraph_rtl(p)
    for run in p.runs:
        run.font.name = "Tahoma"
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    set_paragraph_rtl(p)
    run = p.add_run(text)
    run.font.name = "Tahoma"
    run.font.size = Pt(11)
    run.bold = bold


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph_rtl(p)
        run = p.add_run(item)
        run.font.name = "Tahoma"
        run.font.size = Pt(11)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            set_paragraph_rtl(p)
            for run in p.runs:
                run.bold = True
                run.font.name = "Tahoma"
                run.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, cell_text in enumerate(row):
            row_cells[c_idx].text = cell_text
            for p in row_cells[c_idx].paragraphs:
                set_paragraph_rtl(p)
                for run in p.runs:
                    run.font.name = "Tahoma"
                    run.font.size = Pt(10)
    doc.add_paragraph()


def main() -> None:
    out = Path(__file__).resolve().parents[1] / "docs" / "management" / "regional-manager-sales-targets-proposal-fa.docx"
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.page_height = section.page_height
    # RTL section (optional; paragraph bidi is enough for most Word versions)
    sect_pr = section._sectPr
    bidi = OxmlElement("w:bidi")
    sect_pr.append(bidi)

    add_heading(doc, "پیشنهاد سیستم «تارگت فروش» برای مدیران منطقه", 0)
    add_para(doc, "سند برای گفت‌وگو با مدیریت — بدون جزئیات فنی")
    add_para(doc, "پروژه: mainshop (فروشگاه + پنل نمایندگی + مدیران منطقه)")
    doc.add_paragraph()

    add_heading(doc, "۱. چرا این سیستم لازم است؟", 1)
    add_para(doc, "در وضعیت فعلی:")
    add_bullets(
        doc,
        [
            "مدیران منطقه (مدیر فروش) نمایندگان و مناطق جغرافیایی را مدیریت می‌کنند.",
            "نمایندگان پیش‌فاکتور و سفارش می‌سازند و فروش از طریق پرداخت در سیستم ثبت می‌شود.",
            "مدیر کل مناطق را تعریف می‌کند و می‌بیند هر نماینده زیر کدام مدیر است.",
        ],
    )
    add_para(
        doc,
        "آنچه امروز وجود ندارد: تعریف رسمی «هدف فروش» برای هر مدیر منطقه و مقایسهٔ شفاف "
        "«چقدر باید می‌فروخت» در برابر «چقدر فروخته» در یک بازهٔ زمانی (مثلاً هر ماه).",
    )
    add_para(
        doc,
        "هدف کسب‌وکار: شفافیت، پاسخگویی، برنامه‌ریزی ماهانه، و امکان جلسات عملکرد بدون جمع‌زدن دستی در اکسل.",
    )

    add_heading(doc, "۲. این سیستم دقیقاً چه کاری می‌کند؟", 1)
    add_bullets(
        doc,
        [
            "مدیر کل (یا واحد فروش) برای هر مدیر منطقه در یک دوره (مثلاً «مهر ۱۴۰۴») یک عدد هدف تعیین می‌کند — معمولاً مبلغ فروش به تومان.",
            "سیستم خودکار فروش واقعی همان مدیر را در همان دوره محاسبه می‌کند (بر اساس قوانینی که یک‌بار تأیید می‌کنید).",
            "مدیر منطقه در پنل خود پیشرفت را می‌بیند؛ مثلاً «۶۵٪ از هدف — ۳۲۵ میلیون از ۵۰۰ میلیون».",
            "مدیر کل می‌تواند همهٔ مدیران را در یک نگاه مقایسه کند.",
        ],
    )
    add_para(
        doc,
        "توجه: این ابزار اندازه‌گیری و گزارش است؛ سیستم پاداش یا حقوق نیست مگر مدیریت جداگانه تصمیم بگیرد.",
    )

    add_heading(doc, "۳. «فروش» یعنی چه؟ (باید یک‌بار در مدیریت قطعی شود)", 1)
    add_table(
        doc,
        ["موضوع", "پیشنهاد برای مذاکره"],
        [
            ["چه سفارش‌هایی؟", "فقط فروش از مسیر نمایندگی. فروش مستقیم سایت معمولاً به تارگت مدیر منطقه نمی‌خورد (مگر تصمیم دیگر)."],
            ["چه زمانی به حساب می‌آید؟", "وقتی پول واقعاً واریز شده (پرداخت موفق)، نه فقط صدور پیش‌فاکتور."],
            ["جمع چطور؟", "یا هر مبلغ پرداخت‌شده (حتی قسطی)، یا فقط سفارش کاملاً تسویه‌شده — باید انتخاب شود."],
            ["کدام نمایندگان؟", "نمایندگانی که زیر همان مدیر منطقه هستند (طبق منطق فعلی مناطق و پروفایل نماینده)."],
            ["مرجوعی / لغو", "فاز اول: ساده — کم نمی‌شود؛ فاز بعد: امکان کسر مرجوعی."],
        ],
    )
    add_para(
        doc,
        "نکته: اگر تعریف فروش مکتوب نشود، هر ماه بحث «چرا عدد شما با عدد من فرق دارد» پیش می‌آید.",
        bold=True,
    )

    add_heading(doc, "۴. تارگت به «مدیر» است یا به «منطقه»؟", 1)
    add_para(doc, "الف) یک تارگت برای کل کار مدیر (ساده‌تر)", bold=True)
    add_para(doc, "مثال: مدیر مسئول چند استان است → یک هدف برای مهر؛ فروش = جمع همهٔ نمایندگان زیرمجموعه.")
    add_para(doc, "مناسب وقتی ارزیابی شخص مدیر مهم است.")
    add_para(doc, "ب) تارگت جدا برای هر منطقه (دقیق‌تر)", bold=True)
    add_para(doc, "مثال: منطقه «مرکز» و «قم» هر کدام هدف جدا؛ مدیر دو نوار پیشرفت می‌بیند.")
    add_para(doc, "پیشنهاد شروع: مدل الف (تارگت per مدیر، ماه شمسی، یک رقم مبلغ).")

    add_heading(doc, "۵. دورهٔ زمانی", 1)
    add_table(
        doc,
        ["گزینه", "کاربرد"],
        [
            ["ماه شمسی", "رایج؛ هم‌راستا با گزارش فعلی نمایندگان."],
            ["فصل / سال", "اهداف کلان؛ معمولاً فاز دوم."],
        ],
    )

    add_heading(doc, "۶. چه کسی چه کاری می‌کند؟", 1)
    add_table(
        doc,
        ["نقش", "کار"],
        [
            ["مدیر کل / فروش سازمانی", "تعیین و ویرایش تارگت؛ گزارش همهٔ مدیران؛ تأیید سیاست «فروش یعنی چه»."],
            ["مدیر منطقه", "فقط مشاهده تارگت خود و درصد پیشرفت؛ بدون تغییر عدد هدف."],
            ["نماینده", "فاز اول: همان داشبورد فروش خودش؛ تارگت مدیر اختیاری برای فاز بعد."],
        ],
    )

    add_heading(doc, "۷. چه چیزهایی روی صفحه دیده می‌شود؟", 1)
    add_para(doc, "برای مدیر منطقه (داشبورد):", bold=True)
    add_bullets(
        doc,
        [
            "نام ماه جاری، هدف، تحقق، درصد و نوار پیشرفت",
            "فاصله تا هدف یا روزهای مانده به پایان ماه (اختیاری)",
        ],
    )
    add_para(doc, "برای مدیر کل:", bold=True)
    add_bullets(
        doc,
        [
            "لیست مدیران + ماه + هدف + تحقق + درصد",
            "تعریف / ویرایش تارگت و فیلتر ماه",
        ],
    )
    add_para(doc, "فاز اول عمداً ندارد: پیش‌بینی هوش مصنوعی، تارگت per محصول، اتصال خودکار به حقوق.")

    add_heading(doc, "۸. سناریوهای واقعی (برای تست ذهنی در جلسه)", 1)
    add_bullets(
        doc,
        [
            "مدیر تازه منصوب شده: هدف کل ماه یا تارگت تناسبی؟ (پیشنهاد: هدف کل ماه.)",
            "انتقال نماینده بین مناطق: مدیر فعلی نماینده (ساده) یا مدیر زمان پرداخت (دقیق‌تر).",
            "پرداخت قسطی: هر قسط در ماه پرداخت به فروش همان ماه.",
            "بدون تارگت برای یک ماه: پیام «تارگت تعریف نشده».",
            "مدیر دو منطقه با تارگت per مدیر: یک عدد؛ فروش هر دو منطقه جمع می‌شود.",
        ],
    )

    add_heading(doc, "۹. مزایا برای مدیریت", 1)
    add_bullets(
        doc,
        [
            "یک منبع حقیقت برای عملکرد منطقه‌ای",
            "کاهش جمع‌آوری دستی در اکسل",
            "جلسهٔ ماهانه با اعداد ثابت",
            "پایه برای فاز بعد: پاداش، رتبه‌بندی، تارگت نماینده",
        ],
    )

    add_heading(doc, "۱۰. محدودیت‌ها و ریسک‌ها", 1)
    add_bullets(
        doc,
        [
            "بدون تعریف مکتوب فروش، اعتراض به عدد سیستم محتمل است.",
            "لیست سفارش/پرداخت ادمین هنوز لزوماً فقط «منطقهٔ من» فیلتر نمی‌شود؛ برای تارگت و گزارش کافی است.",
            "تارگت بدون پیگیری سازمانی فقط عدد روی صفحه می‌ماند.",
        ],
    )

    add_heading(doc, "۱۱. پیشنهاد فازبندی", 1)
    add_para(doc, "فاز ۱ — حداقل قابل استفاده:", bold=True)
    add_bullets(
        doc,
        [
            "تارگت مبلغ فروش، ماه شمسی، per مدیر منطقه",
            "تعریف فقط توسط مدیر کل",
            "پیشرفت برای مدیر منطقه + جدول مقایسه برای مدیر کل",
            "قانون فروش: پرداخت موفق سفارش‌های نمایندگان زیرمجموعه",
        ],
    )
    add_para(doc, "فاز ۲ (در صورت تأیید):", bold=True)
    add_bullets(
        doc,
        [
            "تارگت per منطقه یا per نماینده",
            "کسر مرجوعی، گزارش فصلی، export، اعلان درصد هدف",
        ],
    )

    add_heading(doc, "۱۲. سوالات برای جلسه با مدیریت", 1)
    add_bullets(
        doc,
        [
            "هدف فقط مبلغ است یا تعداد سفارش / نمایندهٔ فعال هم لازم است؟",
            "تارگت یک عدد برای کل مدیر است یا برای هر منطقه جدا؟",
            "فروش سایت بدون نماینده به هیچ مدیر منطقه‌ای نمی‌خورد — تأیید؟",
            "پرداخت موفق (حتی قسطی) = فروش همان ماه — تأیید؟",
            "مدیر منطقه فقط ببیند؛ فقط مدیر کل تعریف کند — تأیید؟",
            "ماه شمسی واحد رسمی گزارش — تأیید؟",
            "گزارش ماه‌های قبل از فاز ۱ لازم است یا فقط ماه جاری و آینده؟",
            "این اعداد رسمی برای ارزیابی/پاداش می‌شوند یا فعلاً شاخص مدیریتی؟",
        ],
    )

    add_heading(doc, "۱۳. جمع‌بندی برای ارائه شفاهی", 1)
    add_para(
        doc,
        "برای هر مدیر منطقه، هر ماه یک هدف فروش تعریف می‌کنیم و سیستم خودکار فروش واقعی "
        "نمایندگان همان مدیر را از روی پرداخت‌های موفق جمع می‌زند و درصد تحقق را به مدیر "
        "و به مدیریت نشان می‌دهد. مدیر کل تارگت را set می‌کند؛ مدیر منطقه فقط progress می‌بیند. "
        "در فاز اول یک عدد مبلغ per مدیر per ماه شمسی کافی است. قبل از ساخت، مدیریت باید "
        "یک‌بار تعریف «فروش چیست» را تأیید کند.",
    )

    doc.add_paragraph()
    add_para(doc, "— پایان سند —", bold=True)

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
